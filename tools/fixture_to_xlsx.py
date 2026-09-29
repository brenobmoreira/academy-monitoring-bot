#!/usr/bin/env python3
"""Write a JSON fixture/snapshot (see tools/README.md) back to .xlsx for previewing.

Usage:
  python3 tools/fixture_to_xlsx.py IN.json OUT.xlsx

Written: values, formulas ($formula; =ARRAYFORMULA(x) becomes an array formula), merged cells,
basic formats, column widths, row heights, frozen panes, hidden rows/columns, gridlines, tab
colours, hidden sheets, notes, data validations, conditional formats, auto filter, named ranges,
charts (type, position, size, title, ranges - best effort). Not written: protections, tables
(Google exports tables that overlap the auto filter; Excel rejects that), cached formula values.
"""
import datetime as dt
import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference, ScatterChart, Series
from openpyxl.comments import Comment
from openpyxl.formatting.rule import Rule
from openpyxl.packaging.custom import StringProperty
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from sheetjson import PX_PER_POINT, PX_PER_WIDTH_UNIT, parse_a1, quote_sheet  # noqa: E402

BORDER_STYLES = {'SOLID': 'thin', 'SOLID_MEDIUM': 'medium', 'SOLID_THICK': 'thick', 'DASHED': 'dashed',
                 'DOTTED': 'dotted', 'DOUBLE': 'double'}
DV_OPERATORS = {'BETWEEN': 'between', 'NOT_BETWEEN': 'notBetween', 'EQUAL_TO': 'equal',
                'NOT_EQUAL_TO': 'notEqual', 'GREATER_THAN': 'greaterThan', 'LESS_THAN': 'lessThan',
                'GREATER_THAN_OR_EQUAL_TO': 'greaterThanOrEqual', 'LESS_THAN_OR_EQUAL_TO': 'lessThanOrEqual'}
DATE_OPERATORS = {'DATE_BETWEEN': 'between', 'DATE_NOT_BETWEEN': 'notBetween', 'DATE_EQUAL_TO': 'equal',
                  'DATE_AFTER': 'greaterThan', 'DATE_BEFORE': 'lessThan',
                  'DATE_ON_OR_AFTER': 'greaterThanOrEqual', 'DATE_ON_OR_BEFORE': 'lessThanOrEqual'}
CF_TEXT = {'TEXT_CONTAINS': 'containsText', 'TEXT_DOES_NOT_CONTAIN': 'notContainsText',
           'TEXT_STARTS_WITH': 'beginsWith', 'TEXT_ENDS_WITH': 'endsWith'}
CF_TEXT_FORMULA = {
    'containsText': 'NOT(ISERROR(SEARCH("{t}",{c})))', 'notContainsText': 'ISERROR(SEARCH("{t}",{c}))',
    'beginsWith': 'LEFT({c},LEN("{t}"))="{t}"', 'endsWith': 'RIGHT({c},LEN("{t}"))="{t}"',
}


def argb(color):
    return 'FF' + color.lstrip('#').upper()


def decode_value(v):
    if isinstance(v, dict):
        if '$formula' in v:
            f = v['$formula']
            m = re.match(r'^=ARRAYFORMULA\((.*)\)$', f, re.S | re.I)
            return ('array', '=' + m.group(1)) if m else f
        if '$date' in v:
            return dt.datetime.fromisoformat(v['$date'])
        if '$datetime' in v:
            return dt.datetime.fromisoformat(v['$datetime'])
        if '$time' in v:
            return dt.time.fromisoformat(v['$time'])
        if '$duration' in v:
            return dt.timedelta(seconds=v['$duration'])
    return v


def apply_format(cell, fmt):
    if 'background' in fmt:
        cell.fill = PatternFill('solid', fgColor=argb(fmt['background']))
    font = {}
    if 'fontColor' in fmt:
        font['color'] = argb(fmt['fontColor'])
    if fmt.get('fontWeight') == 'bold':
        font['b'] = True
    if fmt.get('fontStyle') == 'italic':
        font['i'] = True
    if 'fontSize' in fmt:
        font['sz'] = fmt['fontSize']
    if 'fontFamily' in fmt:
        font['name'] = fmt['fontFamily']
    if fmt.get('underline'):
        font['u'] = 'single'
    if fmt.get('strikethrough'):
        font['strike'] = True
    if font:
        cell.font = Font(**font)
    al = {}
    if 'horizontalAlignment' in fmt:
        al['horizontal'] = fmt['horizontalAlignment']
    if 'verticalAlignment' in fmt:
        al['vertical'] = {'middle': 'center'}.get(fmt['verticalAlignment'], fmt['verticalAlignment'])
    if fmt.get('wrap'):
        al['wrap_text'] = True
    if al:
        cell.alignment = Alignment(**al)
    if 'numberFormat' in fmt:
        cell.number_format = fmt['numberFormat']
    if 'borders' in fmt:
        sides = {}
        for side, spec in fmt['borders'].items():
            sides[side] = Side(style=BORDER_STYLES.get(spec.get('style'), 'thin'), color=argb(spec.get('color', '#000000')))
        cell.border = Border(**sides)


def each_cell(ref):
    r1, c1, r2, c2 = parse_a1(ref)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            yield r, c


def dv_formula(v):
    if isinstance(v, dict) and '$date' in v:
        d = dt.date.fromisoformat(v['$date'])
        return str((d - dt.date(1899, 12, 30)).days)
    if isinstance(v, str) and v.startswith('='):
        return v[1:]
    return str(v)


def build_dv(rule):
    crit, values = rule['criteria'], rule.get('values', [])
    kw = {'allow_blank': True, 'showErrorMessage': not rule.get('allowInvalid', True)}
    if rule.get('helpText'):
        kw.update(prompt=rule['helpText'], showInputMessage=True)
    if crit == 'VALUE_IN_LIST':
        items = ','.join(str(x).replace('"', '""') for x in values)
        dv = DataValidation(type='list', formula1=f'"{items}"', **kw)
    elif crit == 'VALUE_IN_RANGE':
        dv = DataValidation(type='list', formula1=str(values[0]).lstrip('='), **kw)
    elif crit.startswith('NUMBER_'):
        dv = DataValidation(type='decimal', operator=DV_OPERATORS.get(crit[7:], 'between'),
                            formula1=dv_formula(values[0]) if values else None,
                            formula2=dv_formula(values[1]) if len(values) > 1 else None, **kw)
    elif crit in DATE_OPERATORS:
        dv = DataValidation(type='date', operator=DATE_OPERATORS[crit],
                            formula1=dv_formula(values[0]) if values else None,
                            formula2=dv_formula(values[1]) if len(values) > 1 else None, **kw)
    elif crit == 'DATE_IS_VALID_DATE':
        dv = DataValidation(type='date', operator='greaterThan', formula1='1', **kw)
    elif crit == 'CHECKBOX':
        dv = DataValidation(type='list', formula1='"TRUE,FALSE"', **kw)
    elif crit == 'CUSTOM_FORMULA':
        dv = DataValidation(type='custom', formula1=dv_formula(values[0]), **kw)
    else:
        return None
    if crit in ('VALUE_IN_LIST', 'VALUE_IN_RANGE') and rule.get('showDropdown') is False:
        dv.showDropDown = True
    return dv


def build_cf(rule):
    fmt = rule.get('format', {})
    font = {}
    if 'fontColor' in fmt:
        font['color'] = argb(fmt['fontColor'])
    if fmt.get('bold'):
        font['b'] = True
    if fmt.get('italic'):
        font['i'] = True
    if fmt.get('strikethrough'):
        font['strike'] = True
    fill = PatternFill('solid', bgColor=argb(fmt['background']), fgColor=argb(fmt['background'])) if 'background' in fmt else None
    dxf = DifferentialStyle(font=Font(**font) if font else None, fill=fill)
    cond, values = rule['condition'], rule.get('values', [])
    first = rule['ranges'][0].split(':')[0]
    if cond == 'CUSTOM_FORMULA':
        return Rule(type='expression', dxf=dxf, formula=[dv_formula(values[0])])
    if cond in CF_TEXT:
        t = CF_TEXT[cond]
        text = str(values[0])
        return Rule(type=t, operator=t, text=text, dxf=dxf,
                    formula=[CF_TEXT_FORMULA[t].format(t=text.replace('"', '""'), c=first)])
    if cond in ('TEXT_EQUAL_TO', 'TEXT_NOT_EQUAL_TO'):
        op = 'equal' if cond == 'TEXT_EQUAL_TO' else 'notEqual'
        return Rule(type='cellIs', operator=op, dxf=dxf, formula=['"' + str(v).replace('"', '""') + '"' for v in values])
    if cond.startswith('NUMBER_'):
        return Rule(type='cellIs', operator=DV_OPERATORS.get(cond[7:], 'equal'), dxf=dxf,
                    formula=[dv_formula(v) for v in values])
    if cond == 'CELL_EMPTY':
        return Rule(type='containsBlanks', dxf=dxf, formula=[f'LEN(TRIM({first}))=0'])
    if cond == 'CELL_NOT_EMPTY':
        return Rule(type='notContainsBlanks', dxf=dxf, formula=[f'LEN(TRIM({first}))>0'])
    return None


CHARTS = {'LINE': LineChart, 'AREA': AreaChart, 'PIE': PieChart, 'SCATTER': ScatterChart,
          'COLUMN': BarChart, 'BAR': BarChart}


def ref_parts(ref, default_sheet):
    m = re.match(r"^(?:'((?:[^']|'')+)'|([^!']+))!(.+)$", ref)
    if m:
        return (m.group(1) or '').replace("''", "'") or m.group(2), m.group(3)
    return default_sheet, ref


def build_chart(wb, ws, spec):
    cls = CHARTS.get(spec.get('type'), LineChart)
    chart = cls()
    if cls is BarChart:
        chart.type = 'bar' if spec['type'] == 'BAR' else 'col'
    opts = spec.get('options', {})
    if opts.get('title'):
        chart.title = opts['title']
    if opts.get('width'):
        chart.width = opts['width'] / 96 * 2.54
        chart.height = opts.get('height', 300) / 96 * 2.54
    for ref in spec.get('ranges', []):
        sheet_name, a1ref = ref_parts(ref, ws.title)
        if sheet_name not in wb.sheetnames:
            continue
        r1, c1, r2, c2 = parse_a1(a1ref)
        data = Reference(wb[sheet_name], min_col=c1, min_row=r1, max_col=c2, max_row=r2)
        for col in range(c1, c2 + 1):
            chart.series.append(Series(Reference(wb[sheet_name], min_col=col, min_row=r1, max_row=r2)))
        del data
    pos = spec.get('position', {'row': 1, 'column': 1})
    ws.add_chart(chart, f"{get_column_letter(pos['column'])}{pos['row']}")


def write_sheet(wb, ws, s):
    for r, line in enumerate(s.get('values', []), start=1):
        for c, v in enumerate(line, start=1):
            if v is None:
                continue
            val = decode_value(v)
            cell = ws.cell(r, c)
            if isinstance(val, tuple):
                cell.value = ArrayFormula(cell.coordinate, val[1])
            else:
                cell.value = val
    for group in s.get('formats', []):
        for ref in group['ranges']:
            for r, c in each_cell(ref):
                apply_format(ws.cell(r, c), group['format'])
    for ref in s.get('merges', []):
        ws.merge_cells(ref)
    for coord, text in s.get('notes', {}).items():
        ws[coord].comment = Comment(text, '')
    for letter, px in s.get('columnWidths', {}).items():
        ws.column_dimensions[letter].width = round(px / PX_PER_WIDTH_UNIT, 2)
    for start, end in s.get('hiddenColumns', []):
        for c in range(start, end + 1):
            ws.column_dimensions[get_column_letter(c)].hidden = True
    for start, end, depth in s.get('columnGroups', []):
        for c in range(start, end + 1):
            ws.column_dimensions[get_column_letter(c)].outlineLevel = depth
    for start, end, depth in s.get('rowGroups', []):
        for r in range(start, end + 1):
            ws.row_dimensions[r].outlineLevel = depth
    for start, end, px in s.get('rowHeights', []):
        for r in range(start, end + 1):
            ws.row_dimensions[r].height = round(px / PX_PER_POINT, 2)
    for start, end in s.get('hiddenRows', []):
        for r in range(start, end + 1):
            ws.row_dimensions[r].hidden = True
    fr, fc = s.get('frozenRows', 0), s.get('frozenColumns', 0)
    if fr or fc:
        ws.freeze_panes = f'{get_column_letter(fc + 1)}{fr + 1}'
    ws.sheet_view.showGridLines = not s.get('hiddenGridlines', False)
    if s.get('tabColor'):
        ws.sheet_properties.tabColor = argb(s['tabColor'])
    if s.get('hidden'):
        ws.sheet_state = 'hidden'
    for rule in s.get('dataValidations', []):
        dv = build_dv(rule)
        if dv is None:
            continue
        dv.sqref = ' '.join(rule['ranges'])
        ws.add_data_validation(dv)
    for rule in s.get('conditionalFormats', []):
        cf = build_cf(rule)
        if cf is not None:
            ws.conditional_formatting.add(' '.join(rule['ranges']), cf)
    if s.get('filter'):
        ws.auto_filter.ref = s['filter']


def convert(data, out_path):
    wb = Workbook()
    wb.remove(wb.active)
    for i, s in enumerate(data['sheets']):
        write_sheet(wb, wb.create_sheet(s['name']), s)
        # xlsx has no grid size; keep it in a custom document property read by xlsx_to_fixture.py.
        if s.get('maxRows') and s.get('maxColumns'):
            wb.custom_doc_props.append(StringProperty(name=f'sheet4.grid.{i}', value=f"{s['maxRows']}x{s['maxColumns']}"))
    for s in data['sheets']:
        for spec in s.get('charts', []):
            build_chart(wb, wb[s['name']], spec)
    for nr in data.get('namedRanges', []):
        r1, c1, r2, c2 = parse_a1(nr['range'])
        ref = f"{quote_sheet(nr['sheet'])}!${get_column_letter(c1)}${r1}"
        if (r1, c1) != (r2, c2):
            ref += f":${get_column_letter(c2)}${r2}"
        wb.defined_names[nr['name']] = DefinedName(nr['name'], attr_text=ref)
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)


def main(argv):
    if len(argv) != 3:
        print(__doc__)
        return 2
    with open(argv[1], encoding='utf-8') as fh:
        convert(json.load(fh), argv[2])
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv))
