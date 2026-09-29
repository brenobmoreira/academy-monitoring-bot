#!/usr/bin/env python3
"""Convert a Google Sheets .xlsx export into a JSON fixture for the sheet 4.0 fakes.

Usage:
  python3 tools/xlsx_to_fixture.py IN.xlsx OUT.json
  python3 tools/xlsx_to_fixture.py --all      # regenerates apps/sheet4/test/fixtures/*_3_0.json

The output shape is documented in tools/README.md. Output is deterministic.
"""
import datetime as dt
import json
import os
import re
import sys
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from sheetjson import (  # noqa: E402
    APPS_SCRIPT_DEFAULTS, PX_PER_POINT, PX_PER_WIDTH_UNIT, cells_to_ranges, col_index, col_letters,
    first_cell_key, norm_number, parse_a1, range_cells, runs,
)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_JOBS = [
    ('sheets/Projeto Breno — Google Sheets.xlsx', 'apps/sheet4/test/fixtures/breno_3_0.json'),
    ('sheets/Projeto Zoio — Google Sheets.xlsx', 'apps/sheet4/test/fixtures/zoio_3_0.json'),
]

THEME_ORDER = ['lt1', 'dk1', 'lt2', 'dk2', 'accent1', 'accent2', 'accent3', 'accent4', 'accent5',
               'accent6', 'hlink', 'folHlink']

BORDER_STYLES = {
    'thin': 'SOLID', 'hair': 'SOLID', 'medium': 'SOLID_MEDIUM', 'thick': 'SOLID_THICK',
    'dashed': 'DASHED', 'mediumDashed': 'DASHED', 'dotted': 'DOTTED', 'double': 'DOUBLE',
}

CHART_TYPES = {
    'LineChart': 'LINE', 'LineChart3D': 'LINE', 'AreaChart': 'AREA', 'AreaChart3D': 'AREA',
    'PieChart': 'PIE', 'PieChart3D': 'PIE', 'DoughnutChart': 'PIE', 'ScatterChart': 'SCATTER',
    'BubbleChart': 'BUBBLE',
}

DV_OPERATORS = {
    None: 'BETWEEN', 'between': 'BETWEEN', 'notBetween': 'NOT_BETWEEN', 'equal': 'EQUAL_TO',
    'notEqual': 'NOT_EQUAL_TO', 'greaterThan': 'GREATER_THAN', 'lessThan': 'LESS_THAN',
    'greaterThanOrEqual': 'GREATER_THAN_OR_EQUAL_TO', 'lessThanOrEqual': 'LESS_THAN_OR_EQUAL_TO',
}
DATE_OPERATORS = {
    None: 'DATE_BETWEEN', 'between': 'DATE_BETWEEN', 'notBetween': 'DATE_NOT_BETWEEN',
    'equal': 'DATE_EQUAL_TO', 'greaterThan': 'DATE_AFTER', 'lessThan': 'DATE_BEFORE',
    'greaterThanOrEqual': 'DATE_ON_OR_AFTER', 'lessThanOrEqual': 'DATE_ON_OR_BEFORE',
}
CF_OPERATORS = {
    'between': 'NUMBER_BETWEEN', 'notBetween': 'NUMBER_NOT_BETWEEN', 'equal': 'NUMBER_EQUAL_TO',
    'notEqual': 'NUMBER_NOT_EQUAL_TO', 'greaterThan': 'NUMBER_GREATER_THAN',
    'lessThan': 'NUMBER_LESS_THAN', 'greaterThanOrEqual': 'NUMBER_GREATER_THAN_OR_EQUAL_TO',
    'lessThanOrEqual': 'NUMBER_LESS_THAN_OR_EQUAL_TO',
}
CF_TEXT = {
    'containsText': 'TEXT_CONTAINS', 'notContainsText': 'TEXT_DOES_NOT_CONTAIN',
    'beginsWith': 'TEXT_STARTS_WITH', 'endsWith': 'TEXT_ENDS_WITH',
}


class Theme:
    def __init__(self, wb):
        self.colors = {}
        raw = getattr(wb, 'loaded_theme', None)
        if not raw:
            return
        ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
        scheme = ET.fromstring(raw).find('.//a:clrScheme', ns)
        if scheme is None:
            return
        for child in scheme:
            tag = child.tag.split('}')[1]
            srgb = child.find('a:srgbClr', ns)
            sysc = child.find('a:sysClr', ns)
            val = srgb.get('val') if srgb is not None else (sysc.get('lastClr') if sysc is not None else None)
            if val:
                self.colors[tag] = '#' + val.lower()

    def resolve(self, color):
        """openpyxl Color -> '#rrggbb' or None. Theme tints are ignored (not used by these exports)."""
        if color is None:
            return None
        if color.type == 'rgb' and isinstance(color.rgb, str):
            rgb = color.rgb
            if len(rgb) == 8:
                if rgb[:2] == '00' and rgb[2:] == '000000':
                    return None  # openpyxl's "no colour" default
                rgb = rgb[2:]
            return '#' + rgb.lower()
        if color.type == 'theme' and color.theme is not None and color.theme < len(THEME_ORDER):
            return self.colors.get(THEME_ORDER[color.theme])
        return None


def encode_value(v):
    if v is None:
        return None
    if isinstance(v, ArrayFormula):
        text = v.text if v.text.startswith('=') else '=' + v.text
        return {'$formula': '=ARRAYFORMULA(' + text[1:] + ')'}
    if isinstance(v, str):
        if v.startswith('='):
            return {'$formula': v}
        return v
    if isinstance(v, bool):
        return v
    if isinstance(v, dt.datetime):
        if v.time() == dt.time(0, 0):
            return {'$date': v.date().isoformat()}
        return {'$datetime': v.replace(microsecond=0).isoformat()}
    if isinstance(v, dt.date):
        return {'$date': v.isoformat()}
    if isinstance(v, dt.time):
        return {'$time': v.replace(microsecond=0).isoformat()}
    if isinstance(v, dt.timedelta):
        return {'$duration': v.total_seconds()}
    if isinstance(v, (int, float)):
        return norm_number(v)
    return str(v)


def encode_cached(v):
    enc = encode_value(v)
    if isinstance(enc, dict) and '$formula' in enc:
        return None
    return enc


def cell_format(cell, theme):
    fmt = {}
    fill = cell.fill
    if fill is not None and fill.fill_type == 'solid':
        bg = theme.resolve(fill.fgColor) or theme.resolve(fill.bgColor)
        if bg and bg != APPS_SCRIPT_DEFAULTS['background']:
            fmt['background'] = bg
    font = cell.font
    if font is not None:
        color = theme.resolve(font.color) if font.color is not None else None
        if color and color != APPS_SCRIPT_DEFAULTS['fontColor']:
            fmt['fontColor'] = color
        if font.b:
            fmt['fontWeight'] = 'bold'
        if font.i:
            fmt['fontStyle'] = 'italic'
        if font.sz is not None and norm_number(float(font.sz)) != APPS_SCRIPT_DEFAULTS['fontSize']:
            # The workbook default (Carlito/Calibri 11) marks cells nobody formatted.
            if not (font.name in ('Carlito', 'Calibri') and float(font.sz) == 11):
                fmt['fontSize'] = norm_number(float(font.sz))
        if font.name and font.name not in ('Arial', 'Carlito', 'Calibri'):
            fmt['fontFamily'] = font.name
        if font.u:
            fmt['underline'] = True
        if font.strike:
            fmt['strikethrough'] = True
    al = cell.alignment
    if al is not None:
        h = {'left': 'left', 'center': 'center', 'centerContinuous': 'center', 'right': 'right',
             'justify': 'left', 'fill': 'left', 'distributed': 'center'}.get(al.horizontal)
        if h:
            fmt['horizontalAlignment'] = h
        v = {'top': 'top', 'center': 'middle', 'justify': 'middle', 'distributed': 'middle'}.get(al.vertical)
        if v:
            fmt['verticalAlignment'] = v
        if al.wrap_text:
            fmt['wrap'] = True
    nf = cell.number_format
    if nf and nf != 'General':
        fmt['numberFormat'] = nf
    b = cell.border
    if b is not None:
        borders = {}
        for side in ('top', 'left', 'bottom', 'right'):
            s = getattr(b, side)
            if s is not None and s.style:
                borders[side] = {'style': BORDER_STYLES.get(s.style, 'SOLID'),
                                 'color': theme.resolve(s.color) or '#000000'}
        if borders:
            fmt['borders'] = borders
    return fmt


def grouped(cell_map, key_name, value_name):
    """{(r,c): value} -> [{ranges, value}] canonical grouping (by JSON-equal value)."""
    groups = {}
    for rc, val in cell_map.items():
        k = json.dumps(val, sort_keys=True, ensure_ascii=False)
        groups.setdefault(k, (val, []))[1].append(rc)
    out = [{key_name: cells_to_ranges(cells), value_name: val} for val, cells in groups.values()]
    out.sort(key=lambda g: first_cell_key(g[key_name]))
    return out


def list_values(formula):
    inner = formula[1:-1].replace('""', '"')
    return [norm_number(_maybe_number(x)) for x in inner.split(',')] if inner else []


def _maybe_number(s):
    try:
        return float(s) if re.match(r'^-?\d+(\.\d+)?$', s.strip()) else s
    except ValueError:
        return s


def dv_value(s):
    if s is None:
        return None
    try:
        return norm_number(float(s))
    except ValueError:
        return '=' + s if not s.startswith('=') else s


def serial_to_date(s):
    try:
        d = dt.date(1899, 12, 30) + dt.timedelta(days=int(float(s)))
        return {'$date': d.isoformat()}
    except ValueError:
        return '=' + s


def encode_dv(dv):
    rule = {}
    f1, f2 = dv.formula1, dv.formula2
    if dv.type == 'list':
        if f1 and f1.startswith('"'):
            rule['criteria'] = 'VALUE_IN_LIST'
            rule['values'] = list_values(f1)
        else:
            rule['criteria'] = 'VALUE_IN_RANGE'
            rule['values'] = [f1]
        if dv.showDropDown:  # xlsx flag is inverted: true hides the arrow
            rule['showDropdown'] = False
    elif dv.type in ('decimal', 'whole'):
        rule['criteria'] = 'NUMBER_' + DV_OPERATORS.get(dv.operator, 'BETWEEN')
        rule['values'] = [dv_value(x) for x in (f1, f2) if x is not None]
    elif dv.type == 'date':
        rule['criteria'] = DATE_OPERATORS.get(dv.operator, 'DATE_BETWEEN')
        rule['values'] = [serial_to_date(x) for x in (f1, f2) if x is not None]
    elif dv.type == 'custom':
        rule['criteria'] = 'CUSTOM_FORMULA'
        rule['values'] = ['=' + f1 if f1 and not f1.startswith('=') else f1]
    else:
        rule['criteria'] = 'XLSX_' + str(dv.type).upper()
        rule['values'] = [x for x in (f1, f2) if x is not None]
    rule['allowInvalid'] = not bool(dv.showErrorMessage)
    if dv.prompt:
        rule['helpText'] = dv.prompt
    return rule


def encode_cf(ws, theme):
    out = []
    for cf in ws.conditional_formatting:
        ranges = sorted(str(cf.sqref).split(), key=lambda r: parse_a1(r)[:2])
        for rule in sorted(cf.rules, key=lambda r: r.priority or 0):
            entry = {'ranges': ranges}
            formulas = list(rule.formula or [])
            if rule.type == 'expression':
                entry['condition'] = 'CUSTOM_FORMULA'
                entry['values'] = ['=' + formulas[0]] if formulas else []
            elif rule.type == 'cellIs':
                vals = [dv_value(f) for f in formulas]
                if all(isinstance(v, str) and v.startswith('="') for v in vals) and rule.operator in ('equal', 'notEqual'):
                    entry['condition'] = 'TEXT_EQUAL_TO' if rule.operator == 'equal' else 'TEXT_NOT_EQUAL_TO'
                    vals = [v[2:-1] for v in vals]
                else:
                    entry['condition'] = CF_OPERATORS.get(rule.operator, 'NUMBER_' + str(rule.operator).upper())
                entry['values'] = vals
            elif rule.type in CF_TEXT:
                entry['condition'] = CF_TEXT[rule.type]
                entry['values'] = [rule.text]
            elif rule.type == 'containsBlanks':
                entry['condition'] = 'CELL_EMPTY'
                entry['values'] = []
            elif rule.type == 'notContainsBlanks':
                entry['condition'] = 'CELL_NOT_EMPTY'
                entry['values'] = []
            else:
                entry['condition'] = 'XLSX_' + str(rule.type).upper()
                entry['values'] = formulas
            fmt = {}
            dxf = rule.dxf
            if dxf is not None:
                if dxf.fill is not None and dxf.fill.fill_type == 'solid':
                    bg = theme.resolve(dxf.fill.bgColor) or theme.resolve(dxf.fill.fgColor)
                    if bg:
                        fmt['background'] = bg
                if dxf.font is not None:
                    if dxf.font.color is not None and theme.resolve(dxf.font.color):
                        fmt['fontColor'] = theme.resolve(dxf.font.color)
                    if dxf.font.b:
                        fmt['bold'] = True
                    if dxf.font.i:
                        fmt['italic'] = True
                    if dxf.font.strike:
                        fmt['strikethrough'] = True
            entry['format'] = fmt
            out.append(entry)
    return out


def rich_text(title):
    if title is None or title.tx is None or title.tx.rich is None:
        return None
    parts = []
    for p in title.tx.rich.p or []:
        for r in p.r or []:
            parts.append(r.t or '')
    text = ''.join(parts)
    return text or None


def series_refs(chart):
    refs = []
    for s in getattr(chart, 'series', []) or []:
        for part in (getattr(s, 'cat', None), getattr(s, 'xVal', None)):
            if part is not None:
                ref = part.numRef or part.strRef
                if ref is not None and ref.f:
                    refs.append(ref.f)
        for part in (getattr(s, 'val', None), getattr(s, 'yVal', None)):
            if part is not None and part.numRef is not None and part.numRef.f:
                refs.append(part.numRef.f)
    seen = []
    for r in refs:
        if r not in seen:
            seen.append(r)
    return seen


def encode_chart(chart):
    name = type(chart).__name__
    ctype = CHART_TYPES.get(name)
    if name in ('BarChart', 'BarChart3D'):
        ctype = 'BAR' if chart.barDir == 'bar' else 'COLUMN'
    anchor = chart.anchor
    frm = anchor._from
    entry = {
        'type': ctype or name.upper(),
        'position': {'row': frm.row + 1, 'column': frm.col + 1,
                     'offsetX': round((frm.colOff or 0) / 9525), 'offsetY': round((frm.rowOff or 0) / 9525)},
        'ranges': [r.replace('$', '') for r in series_refs(chart)],
        'options': {},
    }
    ext = getattr(anchor, 'ext', None)
    if ext is not None and ext.width:
        entry['options']['width'] = round(ext.width / 9525)
        entry['options']['height'] = round(ext.height / 9525)
    title = rich_text(chart.title)
    if title:
        entry['options']['title'] = title
    return entry


def sheet_to_json(ws, cached_ws, theme, grid=None):
    values_by_row = {}
    notes = {}
    fmt_cells = {}
    max_r, max_c = grid or (ws.max_row, ws.max_column)
    max_r, max_c = max(max_r, ws.max_row), max(max_c, ws.max_column)
    for row in ws.iter_rows():
        for cell in row:
            r, c = cell.row, cell.column
            enc = encode_value(cell.value)
            if isinstance(enc, dict) and '$formula' in enc:
                cached = encode_cached(cached_ws.cell(r, c).value)
                if cached is not None:
                    enc['$value'] = cached
            if enc is not None and enc != '':
                values_by_row.setdefault(r, {})[c] = enc
            if cell.comment is not None and cell.comment.text:
                notes[cell.coordinate] = cell.comment.text
            if cell.has_style:
                fmt = cell_format(cell, theme)
                if fmt:
                    fmt_cells[(r, c)] = fmt

    values = []
    if values_by_row:
        for r in range(1, max(values_by_row) + 1):
            line = values_by_row.get(r, {})
            values.append([line.get(c) for c in range(1, max(line) + 1)] if line else [])

    col_widths = {}
    hidden_cols = []
    col_groups = []
    row_groups = []
    for key, dim in ws.column_dimensions.items():
        lo, hi = dim.min or col_index(key), dim.max or col_index(key)
        for c in range(lo, hi + 1):
            if dim.customWidth and dim.width:
                col_widths[c] = round(dim.width * PX_PER_WIDTH_UNIT)
            if dim.hidden:
                hidden_cols.append(c)
            if dim.outlineLevel:
                col_groups.append((c, dim.outlineLevel))
            max_c = max(max_c, c)
    row_heights = []
    hidden_rows = []
    for r, dim in ws.row_dimensions.items():
        if dim.customHeight and dim.ht:
            row_heights.append((r, round(dim.ht * PX_PER_POINT)))
        if dim.hidden:
            hidden_rows.append(r)
        if dim.outlineLevel:
            row_groups.append((r, dim.outlineLevel))
        max_r = max(max_r, r)

    frozen_rows = frozen_cols = 0
    if ws.freeze_panes:
        m = re.match(r'([A-Z]+)(\d+)', ws.freeze_panes)
        frozen_cols, frozen_rows = col_index(m.group(1)) - 1, int(m.group(2)) - 1

    dv_cells = {}
    for dv in ws.data_validations.dataValidation:
        rule = encode_dv(dv)
        for ref in str(dv.sqref).split():
            for rc in range_cells(ref):
                dv_cells[rc] = rule

    tab = ws.sheet_properties.tabColor
    out = {
        'name': ws.title,
        'hidden': ws.sheet_state != 'visible',
        'tabColor': theme.resolve(tab) if tab is not None else None,
        'maxRows': max_r,
        'maxColumns': max_c,
        'frozenRows': frozen_rows,
        'frozenColumns': frozen_cols,
        'hiddenGridlines': not ws.sheet_view.showGridLines if ws.sheet_view.showGridLines is not None else False,
        'columnWidths': {col_letters(c): w for c, w in sorted(col_widths.items())},
        'rowHeights': runs(sorted(row_heights)),
        'hiddenColumns': [[a, b] for a, b, _ in runs([(c, True) for c in sorted(set(hidden_cols))])],
        'hiddenRows': [[a, b] for a, b, _ in runs([(r, True) for r in sorted(set(hidden_rows))])],
        'columnGroups': runs(sorted(col_groups)),
        'rowGroups': runs(sorted(row_groups)),
        'values': values,
        'merges': sorted((str(m) for m in ws.merged_cells.ranges), key=lambda x: parse_a1(x)[:2]),
        'notes': dict(sorted(notes.items(), key=lambda kv: parse_a1(kv[0])[:2])),
        'formats': grouped(fmt_cells, 'ranges', 'format'),
        'dataValidations': [dict(ranges=g['ranges'], **g['rule']) for g in grouped(dv_cells, 'ranges', 'rule')],
        'conditionalFormats': encode_cf(ws, theme),
        'protections': [],
        'charts': [encode_chart(ch) for ch in ws._charts],
        'filter': ws.auto_filter.ref.replace('$', '') if ws.auto_filter and ws.auto_filter.ref else None,
        'tables': [{'name': name, 'range': ref} for name, ref in sorted(ws.tables.items())],
    }
    return out


def named_ranges(wb):
    out = []
    for name, dn in sorted(wb.defined_names.items()):
        if name.startswith('_xlnm') or '.wvu.' in name or dn.hidden:
            continue
        m = re.match(r"^(?:'((?:[^']|'')+)'|([^!]+))!(\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?)$", dn.attr_text or '')
        if not m:
            continue
        sheet = (m.group(1) or '').replace("''", "'") or m.group(2)
        out.append({'name': name, 'sheet': sheet, 'range': m.group(3).replace('$', '')})
    return out


def convert(path):
    wb = openpyxl.load_workbook(path)
    cached = openpyxl.load_workbook(path, data_only=True)
    theme = Theme(wb)
    grids = {}  # written by fixture_to_xlsx.py; Google exports carry the grid as row/col dimensions
    for prop in getattr(wb, 'custom_doc_props', None) or []:
        m = re.match(r'^sheet4\.grid\.(\d+)$', prop.name or '')
        if m and re.match(r'^\d+x\d+$', str(prop.value)):
            grids[int(m.group(1))] = tuple(int(x) for x in str(prop.value).split('x'))
    return {
        'source': os.path.basename(path),
        'sheets': [sheet_to_json(ws, cached[ws.title], theme, grids.get(i)) for i, ws in enumerate(wb.worksheets)],
        'namedRanges': named_ranges(wb),
    }


def dump(data, out_path):
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    with open(out_path, 'w', encoding='utf-8') as fh:
        json.dump(data, fh, ensure_ascii=False, indent=1)
        fh.write('\n')


def main(argv):
    if argv[1:] == ['--all']:
        for src, dst in DEFAULT_JOBS:
            dump(convert(os.path.join(ROOT, src)), os.path.join(ROOT, dst))
            print(f'{src} -> {dst}')
        return 0
    if len(argv) != 3:
        print(__doc__)
        return 2
    dump(convert(argv[1]), argv[2])
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv))
